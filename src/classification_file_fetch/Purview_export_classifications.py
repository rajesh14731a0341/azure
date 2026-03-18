"""
purview_export_classifications.py
==================================
Exports ALL Purview classifications to Excel.

Two sheets:
  - System Classifications   (MICROSOFT.* built-in)
  - Custom Classifications   (tenant-defined)

Three columns each:
  - Display Name
  - Formal Name
  - Description

USAGE:
  python purview_export_classifications.py

CREDENTIALS:
  Uses same credentials as purview_classifier.py
"""

import os
import sys
import requests
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ------------------------------------------------
# CONFIG — same as purview_classifier.py
# ------------------------------------------------

CLIENT_ID       = "2c98dd46-5ec9-4198-b265-058e18078125"
CLIENT_SECRET   = os.environ.get("PURVIEW_CLIENT_SECRET", "")
TENANT_ID       = "5f9bacc0-ffe8-41f7-8d25-215d55cb0f96"
PURVIEW_ACCOUNT = "finastrapurview"

OUTPUT_FILE     = "purview_classifications.xlsx"

# ------------------------------------------------
# AUTH
# ------------------------------------------------

def get_token():
    print("  [AUTH] Fetching token...")
    r = requests.post(
        f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/token",
        data={
            "grant_type":    "client_credentials",
            "client_id":     CLIENT_ID,
            "client_secret": CLIENT_SECRET,
            "resource":      "https://purview.azure.net"
        },
        timeout=30
    )
    r.raise_for_status()
    print("  [AUTH] Token OK")
    return r.json()["access_token"]

# ------------------------------------------------
# FETCH CLASSIFICATIONS
# ------------------------------------------------

def fetch_classifications(token):
    print("\n  Fetching classifications from Purview...")
    url = f"https://{PURVIEW_ACCOUNT}.purview.azure.com/datamap/api/atlas/v2/types/typedefs?type=classification"
    r   = requests.get(url, headers={
        "Authorization": f"Bearer {token}",
        "Content-Type":  "application/json"
    }, timeout=60)

    if r.status_code != 200:
        print(f"  [ERROR] {r.status_code}: {r.text[:300]}")
        sys.exit(1)

    defs = r.json().get("classificationDefs", [])
    print(f"  Total found -> {len(defs)}")
    return defs

# ------------------------------------------------
# SPLIT INTO SYSTEM vs CUSTOM
# ------------------------------------------------

def split_classifications(defs):
    system_rows = []
    custom_rows = []

    for d in defs:
        formal_name  = d.get("name", "")
        display_name = d.get("displayName") or formal_name
        description  = d.get("description", "") or ""

        row = {
            "Display Name": display_name,
            "Formal Name":  formal_name,
            "Description":  description,
        }

        if formal_name.upper().startswith("MICROSOFT."):
            system_rows.append(row)
        else:
            custom_rows.append(row)

    # Sort alphabetically by Display Name
    system_rows.sort(key=lambda x: x["Display Name"].lower())
    custom_rows.sort(key=lambda x: x["Display Name"].lower())

    print(f"  System classifications : {len(system_rows)}")
    print(f"  Custom classifications : {len(custom_rows)}")
    return system_rows, custom_rows

# ------------------------------------------------
# WRITE EXCEL
# ------------------------------------------------

COLS = ["Display Name", "Formal Name", "Description"]

def style_sheet(ws, rows, header_color):
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    row_fill    = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    alt_fill    = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Header row
    for col_idx in range(1, len(COLS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill      = header_fill
        cell.font      = header_font
        cell.border    = thin
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows — alternating colors
    for row_idx in range(2, len(rows) + 2):
        fill = row_fill if row_idx % 2 == 0 else alt_fill
        for col_idx in range(1, len(COLS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill      = fill
            cell.border    = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Column widths
    widths = {"Display Name": 45, "Formal Name": 65, "Description": 60}
    for col_idx, col_name in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths[col_name]

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"


def write_excel(system_rows, custom_rows):
    print(f"\n  Writing Excel -> {OUTPUT_FILE}")

    df_system = pd.DataFrame(system_rows, columns=COLS) if system_rows else pd.DataFrame(columns=COLS)
    df_custom = pd.DataFrame(custom_rows, columns=COLS) if custom_rows else pd.DataFrame(columns=COLS)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df_system.to_excel(writer, sheet_name="System Classifications", index=False)
        df_custom.to_excel(writer, sheet_name="Custom Classifications", index=False)
        style_sheet(writer.sheets["System Classifications"], system_rows, "1F4E79")  # dark blue
        style_sheet(writer.sheets["Custom Classifications"], custom_rows, "375623")  # dark green

    print(f"  Done -> {OUTPUT_FILE}")

# ------------------------------------------------
# MAIN
# ------------------------------------------------

def main():
    print("=" * 60)
    print("  Purview Classifications Exporter")
    print(f"  Account : {PURVIEW_ACCOUNT}")
    print(f"  Output  : {OUTPUT_FILE}")
    print("=" * 60)

    if not CLIENT_SECRET:
        print("\n  [ERROR] PURVIEW_CLIENT_SECRET environment variable is not set.")
        print("  This script uses the same secret as purview_classifier.py.")
        print("  In GitHub Actions it is injected automatically via workflow env:")
        print("    env:")
        print("      PURVIEW_CLIENT_SECRET: ${{ secrets.PURVIEW_CLIENT_SECRET }}")
        sys.exit(1)

    token                    = get_token()
    defs                     = fetch_classifications(token)
    system_rows, custom_rows = split_classifications(defs)
    write_excel(system_rows, custom_rows)

    print(f"\n{'='*60}")
    print(f"  System : {len(system_rows)} classifications")
    print(f"  Custom : {len(custom_rows)} classifications")
    print(f"  Total  : {len(system_rows) + len(custom_rows)} classifications")
    print(f"  File   : {OUTPUT_FILE}")
    print("=" * 60)

if __name__ == "__main__":
    main()